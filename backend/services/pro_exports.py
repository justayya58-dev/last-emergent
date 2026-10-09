"""Professional PDF (reportlab) + Excel (openpyxl) generators for invoices and the
comprehensive business report. Kept decoupled from the DB: callers pass already-fetched
models. Invoice PDF intentionally mirrors the existing endpoint signature."""

from __future__ import annotations

import base64
from datetime import datetime, timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Image,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from models.auditpro import CompanyProfile, Customer, Invoice, ReportSummary, TaxDocument, Transaction
from services.invoice_exports import amount_words

ACCENT_DEFAULT = "#0f9d6f"
INK = colors.HexColor("#1b2430")
MUTED = colors.HexColor("#5b6472")
LIGHT = colors.HexColor("#eef2f6")


def rupiah(value: float) -> str:
    """Rp 1.234.567 with Indonesian dot grouping."""
    whole = int(round(value or 0))
    sign = "-" if whole < 0 else ""
    return f"{sign}Rp {abs(whole):,.0f}".replace(",", ".")


def _accent(company: CompanyProfile) -> colors.Color:
    try:
        return colors.HexColor(company.accent_color or ACCENT_DEFAULT)
    except Exception:
        return colors.HexColor(ACCENT_DEFAULT)


def _logo_image(company: CompanyProfile, max_w: float, max_h: float) -> Image | None:
    if not company.logo_data_url:
        return None
    try:
        raw = base64.b64decode(company.logo_data_url.split(",", 1)[1])
        from reportlab.lib.utils import ImageReader

        reader = ImageReader(BytesIO(raw))
        iw, ih = reader.getSize()
        ratio = min(max_w / iw, max_h / ih)
        return Image(BytesIO(raw), width=iw * ratio, height=ih * ratio)
    except Exception:
        return None


def _styles() -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("title", parent=base["Title"], fontName="Helvetica-Bold", fontSize=22, textColor=INK, spaceAfter=2),
        "h2": ParagraphStyle("h2", parent=base["Heading2"], fontName="Helvetica-Bold", fontSize=13, textColor=INK, spaceBefore=14, spaceAfter=6),
        "label": ParagraphStyle("label", fontName="Helvetica-Bold", fontSize=7.5, textColor=MUTED, spaceAfter=2),
        "body": ParagraphStyle("body", fontName="Helvetica", fontSize=9, textColor=INK, leading=13),
        "small": ParagraphStyle("small", fontName="Helvetica", fontSize=8, textColor=MUTED, leading=11),
        "right": ParagraphStyle("right", fontName="Helvetica", fontSize=9, textColor=INK, alignment=TA_RIGHT),
        "company": ParagraphStyle("company", fontName="Helvetica-Bold", fontSize=15, textColor=INK, spaceAfter=1),
        "docname": ParagraphStyle("docname", fontName="Helvetica-Bold", fontSize=24, textColor=INK, alignment=TA_RIGHT),
        "metaR": ParagraphStyle("metaR", fontName="Helvetica", fontSize=8.5, textColor=MUTED, alignment=TA_RIGHT, leading=12),
    }


# ----------------------------------------------------------------------------- Invoice PDF


def build_invoice_pdf_pro(invoice: Invoice, company: CompanyProfile, customer: Customer | None) -> bytes:
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=16 * mm, title=f"Invoice {invoice.number}")
    s = _styles()
    accent = _accent(company)
    flow: list = []

    logo = _logo_image(company, 42 * mm, 20 * mm)
    left_cell = [logo] if logo else []
    left_cell += [Paragraph(company.name, s["company"]), Paragraph(f"{company.email} · {company.phone}", s["small"]), Paragraph(company.address or "", s["small"])]
    if company.tax_id:
        left_cell.append(Paragraph(f"NPWP: {company.tax_id}", s["small"]))
    status_label = {"draft": "DRAFT", "pending": "FINAL", "sent": "TERKIRIM", "paid": "LUNAS", "overdue": "JATUH TEMPO"}.get(invoice.status, invoice.status.upper())
    right_cell = [
        Paragraph("INVOICE", s["docname"]),
        Spacer(1, 2),
        Paragraph(f"No. <b>{invoice.number}</b>", s["metaR"]),
        Paragraph(f"Tanggal: {invoice.issue_date}", s["metaR"]),
        Paragraph(f"Jatuh tempo: {invoice.due_date}", s["metaR"]),
        Paragraph(f"Status: <b>{status_label}</b>", s["metaR"]),
    ]
    header = Table([[left_cell, right_cell]], colWidths=[100 * mm, 74 * mm])
    header.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
    flow.append(header)
    flow.append(Spacer(1, 6))
    flow.append(Table([[""]], colWidths=[174 * mm], style=TableStyle([("LINEBELOW", (0, 0), (-1, -1), 2, accent)])))
    flow.append(Spacer(1, 10))

    bill_to = [Paragraph("TAGIH KEPADA", s["label"]), Paragraph(f"<b>{invoice.customer}</b>", s["body"])]
    if customer:
        if customer.email or customer.phone:
            bill_to.append(Paragraph(f"{customer.email or ''} {customer.phone or ''}".strip(), s["small"]))
        if customer.address:
            bill_to.append(Paragraph(customer.address, s["small"]))
        if customer.tax_id:
            bill_to.append(Paragraph(f"NPWP: {customer.tax_id}", s["small"]))
    pay_to = [Paragraph("PEMBAYARAN", s["label"]), Paragraph(f"{company.bank_name or '-'}", s["body"]), Paragraph(f"{company.bank_account or '-'}", s["small"])]
    flow.append(Table([[bill_to, pay_to]], colWidths=[100 * mm, 74 * mm], style=TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0)])))
    flow.append(Spacer(1, 12))

    head = ["#", "Deskripsi", "Qty", "Harga", "Jumlah"]
    rows = [head]
    for idx, item in enumerate(invoice.items, start=1):
        rows.append([str(idx), Paragraph(item.description or "-", s["body"]), f"{item.quantity:g} {item.unit}", rupiah(item.price), rupiah(item.total)])
    if len(rows) == 1:
        rows.append(["1", Paragraph("-", s["body"]), "-", rupiah(0), rupiah(0)])
    table = Table(rows, colWidths=[10 * mm, 86 * mm, 22 * mm, 28 * mm, 28 * mm], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), accent),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("ALIGN", (2, 0), (-1, -1), "RIGHT"),
        ("ALIGN", (0, 0), (0, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, LIGHT]),
        ("LINEBELOW", (0, 0), (-1, -1), 0.4, colors.HexColor("#d5dbe2")),
        ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    flow.append(table)
    flow.append(Spacer(1, 10))

    totals = Table([
        ["Subtotal", rupiah(invoice.subtotal)],
        [f"Pajak ({invoice.tax_rate:g}%)", rupiah(invoice.tax)],
        ["TOTAL", rupiah(invoice.total)],
    ], colWidths=[46 * mm, 38 * mm], hAlign="RIGHT")
    totals.setStyle(TableStyle([
        ("ALIGN", (0, 0), (-1, -1), "RIGHT"), ("FONTSIZE", (0, 0), (-1, -1), 9.5),
        ("TEXTCOLOR", (0, 0), (-1, 1), MUTED),
        ("BACKGROUND", (0, 2), (-1, 2), accent), ("TEXTCOLOR", (0, 2), (-1, 2), colors.white),
        ("FONTNAME", (0, 2), (-1, 2), "Helvetica-Bold"), ("FONTSIZE", (0, 2), (-1, 2), 11),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    flow.append(totals)
    flow.append(Spacer(1, 12))
    flow.append(Paragraph("TERBILANG", s["label"]))
    flow.append(Paragraph(f"<i>{invoice.amount_words or amount_words(invoice.total)}</i>", s["body"]))
    if invoice.notes:
        flow.append(Spacer(1, 10))
        flow.append(Paragraph("CATATAN", s["label"]))
        flow.append(Paragraph(invoice.notes, s["small"]))

    doc.build(flow)
    return buf.getvalue()


# ----------------------------------------------------------------------------- Business report PDF


def _section(title: str, head: list[str], rows: list[list], col_widths: list[float], s: dict, accent: colors.Color, aligns: dict[int, str] | None = None, total_row: bool = False) -> list:
    out = [Paragraph(title, s["h2"])]
    data = [head] + (rows if rows else [["Tidak ada data"] + [""] * (len(head) - 1)])
    table = Table(data, colWidths=col_widths, repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), accent),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7.8),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, LIGHT]),
        ("LINEBELOW", (0, 0), (-1, -1), 0.3, colors.HexColor("#d5dbe2")),
        ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]
    for col, align in (aligns or {}).items():
        style.append(("ALIGN", (col, 0), (col, -1), align))
    if total_row and rows:
        style += [("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"), ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#dCeee6")), ("TEXTCOLOR", (0, -1), (-1, -1), INK)]
    table.setStyle(TableStyle(style))
    out.append(table)
    return out


def build_business_report_pdf(company: CompanyProfile, period_start: str, period_end: str, invoices: list[Invoice], transactions: list[Transaction], report: ReportSummary, tax_documents: list[TaxDocument], generated_at: str) -> bytes:
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm, topMargin=16 * mm, bottomMargin=16 * mm, title="Laporan Bisnis")
    s = _styles()
    accent = _accent(company)
    flow: list = []

    logo = _logo_image(company, 40 * mm, 18 * mm)
    ident = [Paragraph(company.name, s["company"]), Paragraph(company.address or "", s["small"]), Paragraph(f"{company.email} · {company.phone}", s["small"])]
    if company.tax_id:
        ident.append(Paragraph(f"NPWP: {company.tax_id}", s["small"]))
    head_left = ([logo, Spacer(1, 4)] if logo else []) + ident
    head_right = [Paragraph("LAPORAN BISNIS", s["docname"]), Spacer(1, 2), Paragraph(f"Periode: {period_start} s/d {period_end}", s["metaR"]), Paragraph(f"Dibuat: {generated_at[:19].replace('T', ' ')}", s["metaR"])]
    header = Table([[head_left, head_right]], colWidths=[104 * mm, 74 * mm], style=TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
    flow += [header, Spacer(1, 6), Table([[""]], colWidths=[178 * mm], style=TableStyle([("LINEBELOW", (0, 0), (-1, -1), 2, accent)]))]

    # 6. Ringkasan Finansial (executive summary up top)
    net_tax = report.tax_collected - report.tax_paid
    flow += _section("Ringkasan Finansial", ["Metrik", "Nilai"], [
        ["Total Penjualan (disetujui)", rupiah(report.sales_total)],
        ["Total Pengeluaran (disetujui)", rupiah(report.expense_total)],
        ["Laba Bersih", rupiah(report.net_profit)],
        ["Piutang (invoice belum lunas)", rupiah(report.receivables)],
        ["Pajak Keluaran", rupiah(report.tax_collected)],
        ["Pajak Masukan", rupiah(report.tax_paid)],
        ["Estimasi PPN Net", rupiah(net_tax)],
        ["Transaksi Disetujui / Menunggu", f"{report.approved_count} / {report.pending_count}"],
    ], [120 * mm, 58 * mm], s, accent, aligns={1: "RIGHT"})
    flow.append(Paragraph(f"Terbilang laba bersih: <i>{amount_words(report.net_profit)}</i>", s["small"]))

    # 1. Invoice
    inv_rows = [[i.number, i.customer[:28], i.issue_date, i.due_date, i.status, rupiah(i.total)] for i in invoices]
    if invoices:
        inv_rows.append(["", "", "", "", "TOTAL", rupiah(sum(i.total for i in invoices))])
    flow += _section("Invoice", ["Nomor", "Customer", "Tanggal", "Jatuh Tempo", "Status", "Total"], inv_rows, [30 * mm, 46 * mm, 24 * mm, 24 * mm, 24 * mm, 30 * mm], s, accent, aligns={5: "RIGHT"}, total_row=True)

    flow.append(PageBreak())

    # 2. Rekap Transaksi
    tx_rows = [[t.transaction_date, ("Masuk" if t.type == "income" else "Keluar"), (t.counterparty or "-")[:26], (t.category or "-")[:18], t.status, rupiah(t.total)] for t in transactions]
    if transactions:
        tx_rows.append(["", "", "", "", "TOTAL", rupiah(sum(t.total for t in transactions))])
    flow += _section("Rekap Transaksi", ["Tanggal", "Tipe", "Pihak", "Kategori", "Status", "Total"], tx_rows, [24 * mm, 18 * mm, 48 * mm, 34 * mm, 24 * mm, 30 * mm], s, accent, aligns={5: "RIGHT"}, total_row=True)

    # 3. Sales Report (approved income)
    sales_tx = [t for t in transactions if t.type == "income" and t.status == "approved"]
    sales_rows = [[t.transaction_date, (t.reference_number or "-")[:20], (t.counterparty or "-")[:26], rupiah(t.subtotal), rupiah(t.tax), rupiah(t.total)] for t in sales_tx]
    if sales_tx:
        sales_rows.append(["", "", "TOTAL", rupiah(sum(t.subtotal for t in sales_tx)), rupiah(sum(t.tax for t in sales_tx)), rupiah(sum(t.total for t in sales_tx))])
    flow += _section("Sales Report (pendapatan disetujui)", ["Tanggal", "Referensi", "Pelanggan", "Subtotal", "Pajak", "Total"], sales_rows, [24 * mm, 30 * mm, 42 * mm, 26 * mm, 26 * mm, 30 * mm], s, accent, aligns={3: "RIGHT", 4: "RIGHT", 5: "RIGHT"}, total_row=True)

    flow.append(PageBreak())

    # 4. Laba Rugi per kategori
    pl_rows = [[c.category, rupiah(c.income), rupiah(c.expense), rupiah(c.income - c.expense)] for c in report.categories]
    pl_rows.append(["LABA BERSIH", rupiah(report.sales_total), rupiah(report.expense_total), rupiah(report.net_profit)])
    flow += _section("Laporan Laba Rugi", ["Kategori", "Pemasukan", "Pengeluaran", "Selisih"], pl_rows, [70 * mm, 36 * mm, 36 * mm, 36 * mm], s, accent, aligns={1: "RIGHT", 2: "RIGHT", 3: "RIGHT"}, total_row=True)

    # 5. Laporan Pajak
    tax_rows = [[d.document_date, ("Keluaran" if d.direction == "output" else "Masukan"), (d.counterparty or "-")[:24], f"{d.tax_rate:g}%", rupiah(d.dpp), rupiah(d.vat)] for d in tax_documents]
    if tax_documents:
        tax_rows.append(["", "", "", "TOTAL", rupiah(sum(d.dpp for d in tax_documents)), rupiah(sum(d.vat for d in tax_documents))])
    flow += _section("Laporan Pajak", ["Tanggal", "Arah", "Lawan Transaksi", "Tarif", "DPP", "PPN"], tax_rows, [24 * mm, 22 * mm, 44 * mm, 18 * mm, 32 * mm, 30 * mm], s, accent, aligns={4: "RIGHT", 5: "RIGHT"}, total_row=True)

    flow.append(Spacer(1, 14))
    flow.append(Paragraph(f"Dokumen dihasilkan otomatis oleh AuditPro AI · {generated_at[:19].replace('T', ' ')} · Semua angka dalam Rupiah.", s["small"]))

    doc.build(flow)
    return buf.getvalue()


# ----------------------------------------------------------------------------- Business report Excel


_RP_FMT = '"Rp" #,##0'
_HDR_FILL = PatternFill("solid", fgColor="0F9D6F")
_TITLE_FONT = Font(bold=True, size=15, color="1B2430")
_HDR_FONT = Font(bold=True, color="FFFFFF", size=10)
_TOTAL_FILL = PatternFill("solid", fgColor="DCEEE6")
_THIN = Side(style="thin", color="D5DBE2")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _sheet(wb: Workbook, title: str, subtitle: str, headers: list[str], rows: list[list], currency_cols: set[int], totals: dict[int, float] | None = None, total_label_col: int = 0) -> None:
    ws: Worksheet = wb.create_sheet(title=title[:31])
    ncols = len(headers)
    last_col = get_column_letter(ncols)
    ws.merge_cells(f"A1:{last_col}1")
    ws["A1"] = title
    ws["A1"].font = _TITLE_FONT
    ws.merge_cells(f"A2:{last_col}2")
    ws["A2"] = subtitle
    ws["A2"].font = Font(italic=True, color="5B6472", size=9)
    header_row = 4
    for ci, name in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=ci, value=name)
        cell.font = _HDR_FONT
        cell.fill = _HDR_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _BORDER
    r = header_row
    for row in rows:
        r += 1
        for ci, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=ci, value=value)
            cell.border = _BORDER
            cell.font = Font(size=9)
            if ci in currency_cols:
                cell.number_format = _RP_FMT
                cell.alignment = Alignment(horizontal="right")
    if totals:
        r += 1
        ws.cell(row=r, column=total_label_col + 1, value="TOTAL").font = Font(bold=True)
        for ci in range(1, ncols + 1):
            cell = ws.cell(row=r, column=ci)
            cell.fill = _TOTAL_FILL
            cell.border = _BORDER
            if ci in totals:
                cell.value = totals[ci]
                cell.number_format = _RP_FMT
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="right")
    # autofilter over header + data (exclude total row for clean filtering)
    data_last = header_row + len(rows)
    ws.auto_filter.ref = f"A{header_row}:{last_col}{max(data_last, header_row)}"
    ws.freeze_panes = f"A{header_row + 1}"
    for ci, name in enumerate(headers, start=1):
        width = max(len(str(name)) + 2, 12)
        for row in rows:
            if ci - 1 < len(row):
                width = max(width, min(len(str(row[ci - 1])) + 2, 42))
        ws.column_dimensions[get_column_letter(ci)].width = width


def build_business_report_xlsx(company: CompanyProfile, period_start: str, period_end: str, invoices: list[Invoice], transactions: list[Transaction], report: ReportSummary, tax_documents: list[TaxDocument], generated_at: str) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    period = f"{company.name} · Periode {period_start} s/d {period_end} · Dibuat {generated_at[:19].replace('T', ' ')}"

    net_tax = report.tax_collected - report.tax_paid
    _sheet(wb, "Ringkasan Finansial", period, ["Metrik", "Nilai"], [
        ["Total Penjualan (disetujui)", report.sales_total],
        ["Total Pengeluaran (disetujui)", report.expense_total],
        ["Laba Bersih", report.net_profit],
        ["Piutang (invoice belum lunas)", report.receivables],
        ["Pajak Keluaran", report.tax_collected],
        ["Pajak Masukan", report.tax_paid],
        ["Estimasi PPN Net", net_tax],
        ["Transaksi Disetujui", report.approved_count],
        ["Transaksi Menunggu", report.pending_count],
    ], currency_cols={2})

    inv_rows = [[i.number, i.customer, i.issue_date, i.due_date, i.status, i.subtotal, i.tax, i.total] for i in invoices]
    _sheet(wb, "Invoice", period, ["Nomor", "Customer", "Tanggal", "Jatuh Tempo", "Status", "Subtotal", "Pajak", "Total"], inv_rows, currency_cols={6, 7, 8}, totals={6: sum(i.subtotal for i in invoices), 7: sum(i.tax for i in invoices), 8: sum(i.total for i in invoices)} if invoices else None)

    tx_rows = [[t.transaction_date, ("Masuk" if t.type == "income" else "Keluar"), t.counterparty, t.reference_number, t.category, t.status, t.subtotal, t.tax, t.total] for t in transactions]
    _sheet(wb, "Rekap Transaksi", period, ["Tanggal", "Tipe", "Pihak", "Referensi", "Kategori", "Status", "Subtotal", "Pajak", "Total"], tx_rows, currency_cols={7, 8, 9}, totals={7: sum(t.subtotal for t in transactions), 8: sum(t.tax for t in transactions), 9: sum(t.total for t in transactions)} if transactions else None)

    sales_tx = [t for t in transactions if t.type == "income" and t.status == "approved"]
    sales_rows = [[t.transaction_date, t.reference_number, t.counterparty, t.subtotal, t.tax, t.total] for t in sales_tx]
    _sheet(wb, "Sales Report", period, ["Tanggal", "Referensi", "Pelanggan", "Subtotal", "Pajak", "Total"], sales_rows, currency_cols={4, 5, 6}, totals={4: sum(t.subtotal for t in sales_tx), 5: sum(t.tax for t in sales_tx), 6: sum(t.total for t in sales_tx)} if sales_tx else None)

    pl_rows = [[c.category, c.income, c.expense, c.income - c.expense] for c in report.categories]
    _sheet(wb, "Laba Rugi", period, ["Kategori", "Pemasukan", "Pengeluaran", "Selisih"], pl_rows, currency_cols={2, 3, 4}, totals={2: report.sales_total, 3: report.expense_total, 4: report.net_profit})

    tax_rows = [[d.document_date, ("Keluaran" if d.direction == "output" else "Masukan"), d.counterparty, d.npwp, f"{d.tax_rate:g}%", d.dpp, d.vat, d.status] for d in tax_documents]
    _sheet(wb, "Laporan Pajak", period, ["Tanggal", "Arah", "Lawan Transaksi", "NPWP", "Tarif", "DPP", "PPN", "Status"], tax_rows, currency_cols={6, 7}, totals={6: sum(d.dpp for d in tax_documents), 7: sum(d.vat for d in tax_documents)} if tax_documents else None)

    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()
