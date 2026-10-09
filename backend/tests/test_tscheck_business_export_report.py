"""Criterion: Business Report PDF/Excel export (multi-section/multi-sheet) + period scoping."""
import io

import openpyxl

EXPECTED_SHEETS = [
    "Ringkasan Finansial",
    "Invoice",
    "Rekap Transaksi",
    "Sales Report",
    "Laba Rugi",
    "Laporan Pajak",
]


def test_business_export_pdf_returns_valid_pdf(client):
    resp = client.get("/auditpro/reports/business-export.pdf")
    assert resp.status_code == 200, resp.text
    assert "pdf" in resp.headers.get("content-type", "").lower()
    assert resp.content[:4] == b"%PDF"
    assert len(resp.content) > 2000, "business report pdf should be non-trivial size"


def test_business_export_pdf_respects_period(client):
    resp = client.get("/auditpro/reports/business-export.pdf", params={"start": "2026-10-01", "end": "2026-10-10"})
    assert resp.status_code == 200, resp.text
    assert resp.content[:4] == b"%PDF"


def test_business_export_xlsx_has_six_named_sheets_with_rupiah_format(client):
    resp = client.get("/auditpro/reports/business-export.xlsx")
    assert resp.status_code == 200, resp.text
    ct = resp.headers.get("content-type", "").lower()
    assert "spreadsheetml" in ct
    assert resp.content[:2] == b"PK"

    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == EXPECTED_SHEETS, wb.sheetnames

    summary_ws = wb["Ringkasan Finansial"]
    assert summary_ws.auto_filter.ref, "summary sheet should have an auto_filter set"

    # Look for at least one cell across sheets using the Rupiah number format.
    found_rupiah_format = False
    for ws in wb.worksheets:
        if found_rupiah_format:
            break
        for row in ws.iter_rows():
            for cell in row:
                if cell.number_format and "Rp" in cell.number_format and "#,##0" in cell.number_format:
                    found_rupiah_format = True
                    break
            if found_rupiah_format:
                break
    assert found_rupiah_format, "expected at least one cell with Rupiah number format '\"Rp\" #,##0'"


def test_business_export_invalid_period_returns_422(client):
    resp = client.get(
        "/auditpro/reports/business-export.pdf",
        params={"start": "2026-10-10", "end": "2026-10-01"},
    )
    assert resp.status_code == 422, resp.text

    resp_bad_format = client.get(
        "/auditpro/reports/business-export.xlsx",
        params={"start": "not-a-date", "end": "2026-10-10"},
    )
    assert resp_bad_format.status_code == 422, resp_bad_format.text
